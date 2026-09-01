"""
Generate the Landing Page Review Form (.docx).

Builds docs/Landing-Page-Review.docx: a plain-language review form covering the
public landing page (index.html) only. Each section of the page gets a feedback
block, followed by a content accuracy check, overall review prompts, a priority
improvements table and a sign-off page.

Usage:
    python scripts/generate-landing-review-doc.py

Requires: python-docx  (pip install python-docx)
"""

from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

# ---------------------------------------------------------------------------
# Configuration
# ---------------------------------------------------------------------------

OUTPUT_PATH = Path(__file__).resolve().parent.parent / "docs" / "Landing-Page-Review.docx"

PAGE_TITLE = "Zeta Beta Mu Fraternity — Landing Page Review"
PAGE_URL = "https://silver-sunflower-cd8320.netlify.app/  (index.html — public landing page)"

BODY_FONT = "Calibri"
HEADING_COLOR = RGBColor(0x0B, 0x3D, 0x2E)
ACCENT_COLOR = RGBColor(0x8A, 0x6D, 0x1F)
MUTED_COLOR = RGBColor(0x55, 0x55, 0x55)

CHECK = "[    ]"

RATING_OPTIONS = ["Good as is", "Needs minor changes", "Needs rework"]
ACTION_OPTIONS = ["Keep", "Change", "Remove"]

# ---------------------------------------------------------------------------
# Landing page section content (source: index.html)
# ---------------------------------------------------------------------------

SECTIONS = [
    (
        "1. Top Navigation Bar & Member Login",
        [
            "ZBM logo with the wordmark \u201cZeta Beta Mu / MEDICAL FRATERNITY\u201d on the left.",
            "Menu links: Home, About, Pillars, Gallery, Vault, Contact.",
            "A light/dark mode switch and a gold \u201cMember Login\u201d button on the right.",
            "The bar sits over the hero at the top, then becomes solid as you scroll; it hides when scrolling down and reappears when scrolling up.",
        ],
    ),
    (
        "2. Mobile Menu (phone / tablet view)",
        [
            "On small screens the menu links are replaced by a hamburger icon.",
            "Tapping it slides in a panel with the logo and the links Home, About, Pillars, Gallery, Contact.",
            "A \u201cMember Login\u201d button sits at the bottom of the panel.",
            "The page behind the panel dims and stops scrolling while the menu is open.",
        ],
    ),
    (
        "3. Hero (first screen)",
        [
            "Badge line: \u201cEST. 1971 \u00b7 UNIVERSITY OF SANTO TOMAS\u201d.",
            "Main title \u201cZeta Beta Mu\u201d with the motto \u201cScholarship. Service. Sociability. Sacrifice. Idealism.\u201d",
            "Short paragraph introducing the fraternity as a distinguished medical brotherhood.",
            "Two buttons: \u201cDiscover Our Legacy\u201d and \u201cOur Five Pillars\u201d.",
            "A large ZBM shield emblem reading \u201cZBM / EST. MCMLXXI\u201d on the right.",
        ],
    ),
    (
        "4. About Us",
        [
            "Small label \u201cDISCOVER OUR STORY\u201d above the heading \u201cAbout Us\u201d.",
            "One long paragraph telling the 1971 founding story and the fraternity's purpose.",
            "A wide heritage photo below the text.",
            "Hovering the photo reveals an \u201cOur Heritage\u201d button that jumps to the gallery.",
        ],
    ),
    (
        "5. The Five Pillars",
        [
            "Label \u201cOUR FOUNDATION\u201d, heading \u201cThe Five Pillars\u201d and the line \u201cThe values that have guided every brother of Zeta Beta Mu since 1971.\u201d",
            "Five cards, each with an icon, a title and a short description: Scholarship, Service, Sociability, Sacrifice, Idealism.",
            "The cards appear one after another as you scroll into the section.",
        ],
    ),
    (
        "6. Numbers / Statistics",
        [
            "Three cards: \u201c55 \u2014 Years of Excellence\u201d, \u201c651 \u2014 Distinguished Alumni\u201d, \u201c1971 \u2014 Founded at UST\u201d.",
            "Each card has an icon and a thin gold bar underneath.",
        ],
    ),
    (
        "7. Villa Maria Integrated School Partnership",
        [
            "Photo of brothers during a community outreach with students, tagged \u201cOngoing Partnership\u201d.",
            "Label \u201cService in Action\u201d and heading \u201cSupporting Villa Maria Integrated School\u201d.",
            "Text explaining support for Aeta students through gift-giving drives, educational supplies and medical outreach.",
            "Two detail tags: \u201cPorac, Pampanga\u201d and \u201cFor Aeta Students\u201d.",
        ],
    ),
    (
        "8. From the Vault (public gallery)",
        [
            "Label \u201cGallery\u201d, heading \u201cFrom the Vault\u201d and a line saying the photos are approved for public viewing.",
            "Eight album cards \u2014 Annual Gala, Induction, Symposium, Reunion, Charity Mission, Founders Day, White Coat Ceremony, Community Health Fair \u2014 each showing a cover photo, category, photo count, venue and year.",
            "A \u201cSee More\u201d button reveals older albums; \u201cShow Less\u201d collapses them again.",
            "Clicking an album opens a photo viewer with next/previous arrows, album details and a \u201cJoin to See More\u201d button.",
            "A locked note reminds visitors that members get the full Historical Vault after logging in.",
        ],
    ),
    (
        "9. Upcoming Events",
        [
            "Label \u201cEvents\u201d, heading \u201cUpcoming Events\u201d and the line \u201cStay connected with what's happening in the Zeta Beta Mu community.\u201d",
            "Three event cards with a date block, title, venue and short description:",
            "MAR 25 \u2014 Annual Gala 2026, The Grand Hotel Ballroom; APR 15 \u2014 New Member Induction, University Great Hall; MAY 10 \u2014 Medical Excellence Conference, Convention Center.",
        ],
    ),
    (
        "10. \u201cBegin Your Legacy\u201d Call to Action",
        [
            "Label \u201cJOIN THE BROTHERHOOD\u201d above the heading \u201cBegin Your Legacy\u201d.",
            "Invitation aimed at UST medical students committed to excellence, service and brotherhood.",
            "Two buttons: \u201cGet in Touch\u201d (sign-up page) and \u201cLearn More\u201d (back to About).",
        ],
    ),
    (
        "11. Footer",
        [
            "Logo, wordmark and a short description of the fraternity and its five pillars.",
            "Facebook and Instagram icons linking to the official pages.",
            "Quick Links: About Us, The Five Pillars, Community Partnerships, Gallery, Member Login.",
            "Contact details: University of Santo Tomas, Espa\u00f1a Blvd., Sampaloc, Manila, Philippines and contact@zetabetamu.org.",
            "Bottom line: \u201c\u00a9 2026 Zeta Beta Mu Medical Fraternity. All rights reserved.\u201d and \u201cEST. 1971 \u00b7 UNIVERSITY OF SANTO TOMAS\u201d.",
        ],
    ),
]

ACCURACY_ITEMS = [
    ("Founding year", "1971, University of Santo Tomas"),
    ("Years of excellence", "55"),
    ("Distinguished alumni", "651"),
    ("The five pillars", "Scholarship, Service, Sociability, Sacrifice, Idealism"),
    ("Partner school", "Villa Maria Integrated School, Porac, Pampanga (Aeta students)"),
    ("Event 1", "Annual Gala 2026 \u2014 Mar 25 \u2014 The Grand Hotel Ballroom"),
    ("Event 2", "New Member Induction \u2014 Apr 15 \u2014 University Great Hall"),
    ("Event 3", "Medical Excellence Conference \u2014 May 10 \u2014 Convention Center"),
    ("Contact email", "contact@zetabetamu.org"),
    ("Postal address", "Espa\u00f1a Blvd., Sampaloc, Manila, Philippines"),
    ("Facebook page", "facebook.com/ZetaBetaMuFraternity"),
    ("Instagram page", "instagram.com/zetabetamufraternity"),
    ("Copyright line", "\u00a9 2026 Zeta Beta Mu Medical Fraternity"),
]

OVERALL_PROMPTS = [
    ("First impression", "What stood out to you in the first few seconds \u2014 good or bad?", 5),
    ("Clarity of message", "Does the page clearly say who we are and what we stand for?", 4),
    ("Confusing or hard to find", "Was anything unclear, hard to read, or hard to locate?", 4),
    ("Missing content", "What should be added to the landing page?", 5),
    ("Content to remove", "What should be taken out or shortened?", 4),
    ("Photos and images", "Are the photos the right ones, and are they good enough quality?", 4),
    ("Wording and tone", "Does the writing sound the way the fraternity should sound?", 4),
    ("Mobile phone experience", "How did the page feel on a phone?", 4),
]


# ---------------------------------------------------------------------------
# Document helpers
# ---------------------------------------------------------------------------

def set_cell_background(cell, hex_color):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    cell._tc.get_or_add_tcPr().append(shd)


def set_cell_borders(cell, color="BFBFBF", size=6):
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), str(size))
        element.set(qn("w:color"), color)
        borders.append(element)
    cell._tc.get_or_add_tcPr().append(borders)


def set_row_height(row, inches):
    tr_pr = row._tr.get_or_add_trPr()
    trHeight = OxmlElement("w:trHeight")
    trHeight.set(qn("w:val"), str(int(inches * 1440)))
    trHeight.set(qn("w:hRule"), "atLeast")
    tr_pr.append(trHeight)


def configure_styles(document):
    normal = document.styles["Normal"]
    normal.font.name = BODY_FONT
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)


def add_spacer(document, points=10):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(points)
    return paragraph


def add_title(document, text, size=26):
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(text)
    run.bold = True
    run.font.size = Pt(size)
    run.font.color.rgb = HEADING_COLOR
    return paragraph


def add_subtitle(document, text, size=12, align=WD_ALIGN_PARAGRAPH.CENTER):
    paragraph = document.add_paragraph()
    paragraph.alignment = align
    run = paragraph.add_run(text)
    run.font.size = Pt(size)
    run.font.color.rgb = MUTED_COLOR
    return paragraph


def add_section_heading(document, text):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(4)
    run = paragraph.add_run(text)
    run.bold = True
    run.font.size = Pt(15)
    run.font.color.rgb = HEADING_COLOR
    return paragraph


def add_label(document, text):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(8)
    paragraph.paragraph_format.space_after = Pt(2)
    run = paragraph.add_run(text.upper())
    run.bold = True
    run.font.size = Pt(9)
    run.font.color.rgb = ACCENT_COLOR
    return paragraph


def add_bullet(document, text):
    paragraph = document.add_paragraph(style="List Bullet")
    paragraph.paragraph_format.space_after = Pt(2)
    run = paragraph.add_run(text)
    run.font.size = Pt(10.5)
    return paragraph


def add_checkbox_row(document, options):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(4)
    for index, option in enumerate(options):
        if index:
            paragraph.add_run("        ")
        box = paragraph.add_run(f"{CHECK}  ")
        box.font.size = Pt(11)
        label = paragraph.add_run(option)
        label.font.size = Pt(10.5)
    return paragraph


def add_writing_box(document, lines=4, caption=None):
    if caption:
        add_label(document, caption)
    table = document.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    cell = table.rows[0].cells[0]
    cell.width = Inches(6.5)
    set_cell_borders(cell)
    set_row_height(table.rows[0], 0.26 * lines)
    cell.paragraphs[0].paragraph_format.space_after = Pt(0)
    for _ in range(max(lines - 1, 0)):
        cell.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_field_line(paragraph, label, underscores=28):
    run = paragraph.add_run(f"{label}: ")
    run.bold = True
    run.font.size = Pt(11)
    line = paragraph.add_run("_" * underscores)
    line.font.size = Pt(11)


def add_page_break(document):
    document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


# ---------------------------------------------------------------------------
# Document sections
# ---------------------------------------------------------------------------

def build_cover(document):
    add_spacer(document, 60)
    add_subtitle(document, "ZETA BETA MU MEDICAL FRATERNITY", size=11)
    add_title(document, "Landing Page Review Form")
    add_subtitle(document, "For review and feedback \u2014 public landing page only")
    add_spacer(document, 24)

    intro = document.add_paragraph()
    intro.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = intro.add_run(
        "This form covers the public landing page \u2014 the page a visitor sees before logging in. "
        "Member-only pages are not part of this review."
    )
    run.font.size = Pt(10.5)
    run.font.color.rgb = MUTED_COLOR

    add_spacer(document, 18)
    add_label(document, "Page reviewed")
    document.add_paragraph(PAGE_URL)

    add_spacer(document, 6)
    add_label(document, "Reviewer details")
    for labels in (
        [("Name", 34), ("Position", 30)],
        [("Date of review", 26), ("Device used (laptop / phone / tablet)", 20)],
        [("Browser used", 30), ("Screen size, if known", 22)],
    ):
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(10)
        for index, (label, width) in enumerate(labels):
            if index:
                paragraph.add_run("        ")
            add_field_line(paragraph, label, width)

    add_spacer(document, 12)
    add_label(document, "How to use this form")
    for step in (
        "Open the landing page and scroll slowly from the top to the very bottom.",
        "Click the buttons, open the menu, open a photo album, and try the links.",
        "View the page once more on a phone if you can.",
        "Fill in one block per section: tick a rating, tick keep / change / remove, then write your comments.",
        "Finish with the Overall Review, list your top improvements in order of priority, and sign at the end.",
    ):
        add_bullet(document, step)

    note = document.add_paragraph()
    note.paragraph_format.space_before = Pt(10)
    run = note.add_run(
        "Please be blunt. Blank space is there to be used \u2014 short notes are more useful than none."
    )
    run.italic = True
    run.font.size = Pt(10)
    run.font.color.rgb = MUTED_COLOR


def build_section_blocks(document):
    add_page_break(document)
    add_title(document, "Part 1 \u2014 Section by Section", size=18)
    add_subtitle(document, "In the order the sections appear on the page")
    add_spacer(document, 10)

    for index, (title, bullets) in enumerate(SECTIONS):
        if index:
            add_page_break(document)

        add_section_heading(document, title)

        add_label(document, "What's on the page now")
        for bullet in bullets:
            add_bullet(document, bullet)

        add_label(document, "Your rating")
        add_checkbox_row(document, RATING_OPTIONS)

        add_label(document, "Keep / change / remove this section")
        add_checkbox_row(document, ACTION_OPTIONS)

        add_writing_box(document, lines=6, caption="Comments and requested changes")


def build_accuracy_check(document):
    add_page_break(document)
    add_title(document, "Part 2 \u2014 Content Accuracy Check", size=18)
    add_subtitle(document, "Please confirm the facts shown on the landing page")
    add_spacer(document, 8)

    table = document.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    table.autofit = False
    widths = [Inches(1.5), Inches(2.7), Inches(0.9), Inches(1.4)]
    headers = ["Item", "What the page says now", "Correct?", "Correction, if any"]

    for cell, header, width in zip(table.rows[0].cells, headers, widths):
        cell.width = width
        set_cell_background(cell, "0B3D2E")
        paragraph = cell.paragraphs[0]
        paragraph.paragraph_format.space_after = Pt(2)
        run = paragraph.add_run(header)
        run.bold = True
        run.font.size = Pt(10)
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    for item, value in ACCURACY_ITEMS:
        row = table.add_row()
        set_row_height(row, 0.34)
        cells = row.cells
        for cell, width in zip(cells, widths):
            cell.width = width
        for cell, text in zip(cells, (item, value, f"{CHECK} Yes   {CHECK} No", "")):
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(2)
            run = paragraph.add_run(text)
            run.font.size = Pt(9.5)

    add_spacer(document, 10)
    add_writing_box(document, lines=4, caption="Other facts that need fixing")


def build_overall_review(document):
    add_page_break(document)
    add_title(document, "Part 3 \u2014 Overall Review", size=18)
    add_subtitle(document, "Your general impression of the landing page as a whole")
    add_spacer(document, 8)

    for index, (title, prompt, lines) in enumerate(OVERALL_PROMPTS):
        if index == 4:
            add_page_break(document)
        add_section_heading(document, title)
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(4)
        run = paragraph.add_run(prompt)
        run.italic = True
        run.font.size = Pt(10.5)
        run.font.color.rgb = MUTED_COLOR
        add_writing_box(document, lines=lines)
        add_spacer(document, 6)

    add_page_break(document)
    add_section_heading(document, "Overall verdict on the landing page")
    add_checkbox_row(
        document,
        ["Excellent", "Good", "Acceptable", "Needs work", "Start over"],
    )
    add_spacer(document, 6)
    add_writing_box(document, lines=5, caption="Anything else you want the team to know")


def build_priority_table(document):
    add_page_break(document)
    add_title(document, "Part 4 \u2014 Priority Improvements", size=18)
    add_subtitle(document, "List the changes you want, most important first")
    add_spacer(document, 8)

    table = document.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    table.autofit = False
    widths = [Inches(0.4), Inches(3.5), Inches(1.3), Inches(1.3)]
    headers = ["#", "Improvement requested", "Which section", "Priority"]

    for cell, header, width in zip(table.rows[0].cells, headers, widths):
        cell.width = width
        set_cell_background(cell, "0B3D2E")
        paragraph = cell.paragraphs[0]
        paragraph.paragraph_format.space_after = Pt(2)
        run = paragraph.add_run(header)
        run.bold = True
        run.font.size = Pt(10)
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    for number in range(1, 11):
        row = table.add_row()
        set_row_height(row, 0.42)
        cells = row.cells
        for cell, width in zip(cells, widths):
            cell.width = width
        number_run = cells[0].paragraphs[0].add_run(str(number))
        number_run.font.size = Pt(10)
        number_run.bold = True
        priority = cells[3].paragraphs[0].add_run("High / Med / Low")
        priority.font.size = Pt(9)
        priority.font.color.rgb = MUTED_COLOR


def build_signoff(document):
    add_page_break(document)
    add_title(document, "Part 5 \u2014 Sign-Off", size=18)
    add_spacer(document, 10)

    add_label(document, "Decision on the landing page")
    for option in (
        "Approved \u2014 proceed to final polish.",
        "Approved with changes \u2014 apply the requested changes above, no further review needed.",
        "Needs another round \u2014 apply the changes and send the page back to me for review.",
    ):
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(8)
        box = paragraph.add_run(f"{CHECK}  ")
        box.font.size = Pt(11)
        label = paragraph.add_run(option)
        label.font.size = Pt(10.5)

    add_spacer(document, 18)
    add_writing_box(document, lines=4, caption="Conditions or deadline")

    add_spacer(document, 28)
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(18)
    add_field_line(paragraph, "Reviewed by", 34)
    paragraph.add_run("        ")
    add_field_line(paragraph, "Date", 20)

    paragraph = document.add_paragraph()
    add_field_line(paragraph, "Signature", 40)

    add_spacer(document, 24)
    footer_note = document.add_paragraph()
    footer_note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer_note.add_run(
        "Zeta Beta Mu Medical Fraternity \u2014 landing page review. "
        "Return this form to the web team once completed."
    )
    run.font.size = Pt(9)
    run.font.color.rgb = MUTED_COLOR


def build_document():
    document = Document()
    configure_styles(document)

    section = document.sections[0]
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)

    footer_paragraph = section.footer.paragraphs[0]
    footer_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run = footer_paragraph.add_run(f"{PAGE_TITLE}  \u00b7  Landing page only")
    footer_run.font.size = Pt(8)
    footer_run.font.color.rgb = MUTED_COLOR

    build_cover(document)
    build_section_blocks(document)
    build_accuracy_check(document)
    build_overall_review(document)
    build_priority_table(document)
    build_signoff(document)

    return document


def main():
    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    document = build_document()
    document.save(OUTPUT_PATH)
    print(f"Wrote {OUTPUT_PATH}")


if __name__ == "__main__":
    main()
